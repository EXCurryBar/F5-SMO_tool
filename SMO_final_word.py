import paramiko
import time
import os
import utils
import multiprocessing as mp
import logging
import pandas as pd
from scp import SCPClient
from docx import Document
from docx.shared import Cm

logging.basicConfig(level=logging.WARNING,
                    format="%(asctime)s %(levelname)s %(message)s",
                    datefmt="%Y-%m-%d %H:%M",
                    handlers=[logging.FileHandler("SMO.log", "w", "utf-8"), ])


def sys_search(con, client):
    try:
        stdin, stdout, stderr = client.exec_command(con)
        for line in stdout.readlines():
            sys_text = line.split(" ")
            return sys_text
    except:
        sys_text = 'fail'
        print('fail'+con)
        return sys_text


def sys_compare(con, client):
    try:
        stdin, stdout, stderr = client.exec_command(con)
        for line in stdout.readlines():
            if "none" in line:
                return "N/A"
            else:
                return "OK"
    except:
        return "fail"
        print('fail'+con)


def charge(ex):
    ex = ex[-1]
    ex = ex[:-1]
    return ex


def words(ti, da, num):
    global word_nn, doc, t0
    print(ti+": "+da)
    if word_nn == 1:
        t0.cell(num, word_nn).text = da
    elif word_nn == 3:
        t0.cell(num, word_nn).text = da
    elif word_nn == 5:
        t0.cell(num, word_nn).text = da
    else:
        t0.cell(num, 6).text = da
    while True:
        try:
            doc.save('SMO_v1.docx')
            break
        except:
            pa = str(input("請關閉文件再試一次: "))


def paste(sys_uptime, sys_host, sys_sn, sys_ver, sys_Edit, sys_cpu, sys_tmm_mem, sys_con, sys_ccon, sys_ith, sys_time, lt, sys_ntp, sys_snmp, sys_ha):
    sys_uptime = sys_uptime[-2]
    sys_host = charge(sys_host)
    sys_sn = charge(sys_sn)
    sys_ver = charge(sys_ver)
    sys_cpu = sys_cpu
    sys_tmm_mem = sys_tmm_mem
    sys_ccon = charge(sys_ccon)
    sys_ith = charge(sys_ith)
    sys_con = sys_con[:-1]
    sys_ha = sys_ha[:-1]
    sys_Edit = str(sys_Edit[-3])+" "+str(sys_Edit[-2])+" "+str(sys_Edit[-1])
    sys_Edit = sys_Edit[:-1]
    words("hostname", sys_host, 0)
    words("S/N", sys_sn, 1)
    words("uptime", sys_uptime+" days", 2)
    words("Memory", sys_tmm_mem, 3)
    words("CPU", sys_cpu, 4)
    words("Active Connections", sys_con, 5)
    words("New Connections", sys_ccon+"/sec", 6)
    words("Throughput", sys_ith+"(bits/sec)", 7)
    words("NTP", sys_ntp, 9)
    words("SNMP", sys_snmp, 10)
    sys_time_t = sys_time.split(":")
    print("F5 TIME: "+sys_time)
    lt = str(lt[-2])
    lt_t = lt.split(":")
    print("local time: "+lt)
    if lt_t[0] == sys_time_t[0]:
        if lt_t[1] == sys_time_t[1]:
            text = "same"
            words("Time", text, 13)
        elif lt_t[1] > sys_time_t[1]:
            text = "慢"+str(int(lt_t[1])-int(sys_time_t[1]))+"分鐘"
            words("Time", text, 13)
        else:
            text = "快"+str(int(sys_time_t[1])-int(lt_t[1]))+"分鐘"
            words("Time", text, 13)
    elif lt_t[0] > sys_time_t[0]:
        if lt_t[1] == sys_time_t[1]:
            text = "慢"+str(int(lt_t[0])-int(sys_time_t[0]))+"小時"
            words("Time", text, 13)
        elif lt_t[1] > sys_time_t[1]:
            text = "慢"+str(int(lt_t[0])-int(sys_time_t[0])) + \
                "小時"+str(int(lt_t[1])-int(sys_time_t[1]))+"分鐘"
            words("Time", text, 13)
        else:
            text = "慢" + \
                str(((int(lt_t[0])-int(sys_time_t[0]))*60) -
                    int(sys_time_t[1])+int(lt_t[1]))+"分鐘"
            words("Time", text, 13)
    else:
        if lt_t[1] == sys_time_t[1]:
            text = "快"+str(int(sys_time_t[0])-int(lt_t[0]))+"小時"
            words("Time", text, 13)
        elif sys_time_t[1] > lt_t[1]:
            text = "快"+str(int(sys_time_t[0])-int(lt_t[0])) + \
                "小時"+str(int(sys_time_t[1])-int(lt_t[1]))+"分鐘"
            words("Time", text, 13)
        else:
            text = "快" + \
                str(((int(sys_time_t[0])-int(lt_t[0]))*60) -
                    int(lt_t[1])+int(sys_time_t[1]))+"分鐘"
            words("Time", text, 13)
    words("HA ststus", sys_ha, 15)
    words("Version", sys_ver+" "+sys_Edit, 16)


def ver_13(IP, user, passwd):
    decide_num = 0
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(IP, 22, username=user, password=passwd, timeout=10)
    sys_uptime = sys_search("tmsh show sys service big3d", client)
    sys_host = sys_search(
        'tmsh list sys global-settings hostname | grep "hostname"', client)
    sys_sn = sys_search(
        'tmsh show sys hardware | grep "Chassis Serial"', client)
    sys_ver = sys_search('tmsh show sys version | grep " Version"', client)
    try:
        sys_data = utils.get_data(IP, user, passwd)
    except:
        print("連線速度過慢 :" + IP + "跳過cpu, mem")
        logging.error("連線速度過慢 :" + IP + "跳過cpu, mem")
    sys_tmm_mem = sys_data[0][1]
    sys_cpu = sys_data[1][1]
    # sys_cpu = sys_search(
    #     'tmsh show sys performance system historical | grep "Utilization"', client)
    # sys_tmm_mem = sys_search(
    #     'tmsh show sys performance system historical | grep "TMM Memory Used"', client)
    sys_ccon = sys_search(
        'tmsh show sys performance connections historical | grep "Client Connections"', client)
    sys_Edit = sys_search('tmsh show sys version | grep "Edition"', client)
    sys_ith = sys_search(
        'tmsh show sys performance throughput historical | grep "In"', client)
    sys_ntp = sys_compare('tmsh list sys ntp servers | grep "servers"', client)
    sys_snmp = sys_compare(
        'tmsh list sys snmp allowed-addresses | grep "allowed-addresses"', client)
    stdin, stdout, stderr = client.exec_command(
        'tmsh show cm sync-status | grep "Status"')
    for line in stdout.readlines():
        decide_num += 1
        sys_ha = line.split(" ")
        if decide_num == 2:
            sys_ha = sys_ha[-1]
            decide_num = 0
            break
    stdin, stdout, stderr = client.exec_command(
        'tmsh show sys performance connections historical | grep "Connections"')
    for line in stdout.readlines():
        decide_num += 1
        sys_con = line.split(" ")
        if decide_num == 3:
            sys_con = sys_con[-1]
            decide_num = 0
            break
    stdin, stdout, stderr = client.exec_command('tmsh show sys clock')
    lt = time.ctime()
    lt = lt.split(" ")
    for line in stdout.readlines():
        decide_num += 1
        sys_time = line.split(" ")
        if decide_num == 4:
            sys_time = str(sys_time[-3])
            break
    client.close()
    paste(sys_uptime, sys_host, sys_sn, sys_ver, sys_Edit, sys_cpu, sys_tmm_mem,
          sys_con, sys_ccon, sys_ith, sys_time, lt, sys_ntp, sys_snmp, sys_ha)


def Compare_final(ex, IP, user, passwd):
    if ex == 11:
        print('程式目前適用於13版')
        logging.warning('程式目前適用於13版' + IP)
    else:
        ver_13(IP, user, passwd)


def Compare_ver(IP, user, passwd):
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    decide_num = 1
    while decide_num:
        try:
            client.connect(IP, 22, username=user, password=passwd, timeout=10)
            stdin, stdout, stderr = client.exec_command(
                'tmsh show sys version | grep " Version"')
            for line in stdout.readlines():
                sys_com_ver = line.split(" ")
                sys_com_ver = str(sys_com_ver[-1])
                sys_com_ver = sys_com_ver.split(".")
            client.close()
            Compare_num = 0
        except:
            print("IP錯誤或帳密錯誤，請檢查帳密")
            logging.error("IP錯誤或帳密錯誤，請檢查帳密:" + IP)
            decide_num = 0
            continue
        if Compare_num == 0:
            Compare_final(int(sys_com_ver[0]), IP, user, passwd)
            decide_num = 0


def main():
    global doc, t0, word_nn
    PATH = os.path.abspath(os.getcwd())
    os.chdir("\\")
    os.system("mkdir qkviews, ucs")
    os.chdir(PATH)
    doc = Document('example.docx')
    t0 = doc.tables[0]
    word_nn = 1
    data = pd.read_excel("SMO_ex.xls").values.tolist()
    for row in data:
        IP, user, passwd = row
        if utils.is_avail(IP, 443):
            print(IP, user, passwd)
            Compare_ver(IP, user, passwd)
            word_nn += 2
        else:
            print("無法連線到 :", IP)
            logging.error("無法連線到 :" + IP)


if __name__ == "__main__":
    main()
